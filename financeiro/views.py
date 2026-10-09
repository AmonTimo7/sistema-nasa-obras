from datetime import date, timedelta
from io import BytesIO

from django.contrib import messages
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Max, Sum
from django.db.models.functions import Coalesce
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST

from compras.models import Fornecedor
from core.models import Perfil
from core.permissoes import categorias_permitidas, exige_papel, papel_de, so_financeiro
from juridico.models import Prestador

from . import consultas, regras
from .forms import (
    AditivoForm,
    ContaForm,
    EmpreitaForm,
    ItemChecklistForm,
    LancamentoForm,
    LocacaoForm,
    MedicaoForm,
    ObraForm,
    PagarContaForm,
)
from .models import (
    FORMAS_PAGAMENTO,
    ZERO,
    Aditivo,
    Categoria,
    ContaAPagar,
    Empreita,
    ItemChecklist,
    Lancamento,
    LocacaoRecorrente,
    Medicao,
    Obra,
    Responsavel,
    Subcategoria,
)

PODE_LANCAR = (Perfil.COMPRAS,)


# ---------------------------------------------------------------------------
# obra em foco
# ---------------------------------------------------------------------------

def obra_atual(request):
    """A obra em foco fica na sessão; ?obra=ID troca."""
    escolhida = request.GET.get("obra")
    if escolhida and escolhida.isdigit() and Obra.objects.filter(pk=escolhida).exists():
        request.session["obra_id"] = int(escolhida)
    obra = None
    if request.session.get("obra_id"):
        obra = Obra.objects.filter(pk=request.session["obra_id"]).first()
    obra = obra or Obra.objects.exclude(situacao=Obra.CONCLUIDA).first() or Obra.objects.first()
    if obra:
        request.session["obra_id"] = obra.pk
    return obra


def _sem_obra(request):
    if papel_de(request.user) == Perfil.SOCIO:
        messages.info(request, "Cadastre a primeira obra para começar.")
        return redirect("financeiro:obra_nova")
    return render(request, "core/sem_modulo.html")


def com_obra(view):
    """Injeta a obra atual; sem obra cadastrada, manda cadastrar."""
    def embrulho(request, *args, **kwargs):
        obra = obra_atual(request)
        if obra is None:
            return _sem_obra(request)
        return view(request, obra, *args, **kwargs)
    embrulho.__name__ = view.__name__
    embrulho.__doc__ = view.__doc__
    return embrulho


def _ctx(request, obra, **extra):
    extra.setdefault("obra", obra)
    extra.setdefault("obras", Obra.objects.all())
    return extra


# ---------------------------------------------------------------------------
# painel
# ---------------------------------------------------------------------------

@so_financeiro
@com_obra
def painel(request, obra):
    hoje = date.today()
    regras.gerar_todas_as_locacoes()
    r = consultas.resumo_obra(obra, hoje)
    por_mes = consultas.gasto_por_mes(obra)
    acum = consultas.acumulado(obra)
    graficos = {
        "meses": [consultas.nome_mes(m, curto=True) for m, _ in por_mes],
        "chaves": [m.strftime("%Y/%m") for m, _ in por_mes],
        "gasto": [float(v) for _, v in por_mes],
        "acum_gasto": [float(a["gasto"]) for a in acum],
        "acum_faturado": [float(a["faturado"]) for a in acum],
    }
    return render(request, "financeiro/painel.html", _ctx(
        request, obra, r=r, hoje=hoje, graficos=graficos,
        categorias=consultas.gasto_por_categoria(obra),
        favorecidos=consultas.maiores_favorecidos(obra, 7),
        contas=consultas.contas_resumo(obra, hoje),
        proximas=consultas.proximas_contas(obra, 5, hoje),
        n_pendencias=len(consultas.pendencias(obra)),
        mes_atual=consultas.nome_mes(hoje),
    ))


@so_financeiro
def visao_geral(request):
    regras.gerar_todas_as_locacoes()
    hoje = date.today()
    return render(request, "financeiro/visao_geral.html", {
        "v": consultas.visao_geral(hoje), "obra": obra_atual(request), "obras": Obra.objects.all(),
        "contas": consultas.contas_resumo(None, hoje), "proximas": consultas.proximas_contas(None, 8, hoje),
        "mes_atual": consultas.nome_mes(hoje),
    })


# ---------------------------------------------------------------------------
# lançamentos
# ---------------------------------------------------------------------------

def _pode_categoria(request, codigo):
    permitidas = categorias_permitidas(request.user)
    return permitidas is None or codigo in permitidas


@exige_papel(*PODE_LANCAR)
@com_obra
def lancamentos(request, obra):
    permitidas = categorias_permitidas(request.user)
    qs = consultas.lancamentos_filtrados(obra, request.GET, permitidas)
    total = qs.aggregate(t=Coalesce(Sum("valor"), ZERO))["t"]
    pagina = Paginator(qs, 40).get_page(request.GET.get("pagina"))
    for l in pagina:
        l.status = regras.status_pagamento(l)
    filtros = request.GET.copy()
    filtros.pop("pagina", None)
    cats = Categoria.objects.all() if permitidas is None else Categoria.objects.filter(codigo__in=permitidas)
    contexto = _ctx(
        request, obra, pagina=pagina, total=total, quantidade=qs.count(), categorias=cats,
        meses=[(m.strftime("%Y-%m"), consultas.nome_mes(m)) for m in consultas.meses_da_obra(obra)],
        f=request.GET, filtros_qs=filtros.urlencode(), formas=FORMAS_PAGAMENTO,
        filtro_nome=_nome_do_filtro(request.GET),
    )
    if request.headers.get("HX-Request") and request.headers.get("HX-Target") == "tabela":
        return render(request, "financeiro/_tabela_lancamentos.html", contexto)
    return render(request, "financeiro/lancamentos.html", contexto)


def _nome_do_filtro(params):
    for campo, modelo, rotulo in (("prestador", Prestador, "de "), ("fornecedor", Fornecedor, "de "),
                                  ("empreita", Empreita, "da empreita "), ("subcategoria", Subcategoria, "de ")):
        valor = params.get(campo) or ""
        if valor.isdigit():
            obj = modelo.objects.filter(pk=valor).first()
            return f"{rotulo}{obj}" if obj else ""
    return ""


@exige_papel(*PODE_LANCAR)
@com_obra
def lancamento_escolher(request, obra):
    permitidas = categorias_permitidas(request.user)
    cats = Categoria.objects.all() if permitidas is None else Categoria.objects.filter(codigo__in=permitidas)
    return render(request, "financeiro/lancamento_escolher.html", _ctx(request, obra, categorias=cats))


@exige_papel(*PODE_LANCAR)
@com_obra
def lancamento_form(request, obra, pk=None, codigo=None):
    lanc = get_object_or_404(Lancamento, pk=pk, obra=obra) if pk else None
    categoria = lanc.categoria if lanc else get_object_or_404(Categoria, codigo=codigo)
    if not _pode_categoria(request, categoria.codigo):
        raise PermissionDenied
    if lanc and lanc.medicao_id:
        messages.info(request, "Este imposto é gerado pela medição. Edite a medição.")
        return redirect("financeiro:medicao_editar", pk=lanc.medicao_id)
    if lanc and lanc.conta_origem_id:
        messages.info(request, "Este gasto nasceu do pagamento de uma locação mensal. Mude pela conta.")
        return redirect("financeiro:contas")
    inicial = {}
    if not lanc:
        inicial = {"data": date.today()}
        for campo in ("prestador", "empreita", "estorno_de", "subcategoria"):
            if (request.GET.get(campo) or "").isdigit():
                inicial[campo] = request.GET[campo]
        if inicial.get("prestador"):
            p = Prestador.objects.filter(pk=inicial["prestador"]).first()
            if p:
                inicial["valor_diaria"] = p.valor_diaria
    form = LancamentoForm(request.POST or None, request.FILES or None, instance=lanc, initial=inicial,
                          categoria=categoria, obra=obra)
    avisos = []
    if request.method == "POST" and form.is_valid():
        if form.avisos and not request.POST.get("confirmar"):
            avisos = form.avisos
        else:
            with transaction.atomic():
                obj = form.save(commit=False)
                if not lanc:
                    obj.criado_por = request.user
                obj.save()
                regras.sincronizar_contas(obj)
            messages.success(request, f"R$ {obj.valor} lançado em {categoria.nome}.")
            if request.POST.get("e_outro"):
                url = reverse("financeiro:lancamento_novo", args=[categoria.codigo])
                return redirect(url)
            return redirect("financeiro:lancamento", pk=obj.pk)
    saldo_vales = None
    if categoria.codigo == Categoria.MAO_DE_OBRA:
        pid = form["prestador"].value()
        if pid and str(pid).isdigit():
            p = Prestador.objects.filter(pk=pid).first()
            if p:
                saldo_vales = regras.saldo_de_vales(p, obra, ignorar=lanc)
    especificos, comuns = form.ordem()
    return render(request, "financeiro/lancamento_form.html", _ctx(
        request, obra, form=form, l=lanc, categoria=categoria, saldo_vales=saldo_vales, avisos=avisos,
        especificos=[form[n] for n in especificos], comuns=[form[n] for n in comuns],
        fornecedores=Fornecedor.objects.filter(ativo=True),
        diarias={p.pk: f"{p.valor_diaria:.2f}".replace(".", ",") for p in Prestador.objects.filter(ativo=True)},
    ))


@exige_papel(*PODE_LANCAR)
def lancamento_detalhe(request, pk):
    l = get_object_or_404(Lancamento.objects.select_related(
        "categoria", "subcategoria", "prestador", "fornecedor", "empreita", "conta_bancaria", "responsavel", "obra"),
        pk=pk)
    if not _pode_categoria(request, l.categoria.codigo):
        raise PermissionDenied
    return render(request, "financeiro/lancamento_detalhe.html", _ctx(
        request, l.obra, l=l, status=regras.status_pagamento(l), contas=l.contas.all(),
        estornos=l.estornos.all(), parecidos=regras.lancamentos_parecidos(l),
    ))


@exige_papel(*PODE_LANCAR)
@require_POST
def lancamento_excluir(request, pk):
    l = get_object_or_404(Lancamento, pk=pk)
    if not _pode_categoria(request, l.categoria.codigo):
        raise PermissionDenied
    if l.contas.filter(pago_em__isnull=False).exists():
        messages.error(request, "Não dá para excluir: há parcela paga. Faça um estorno.")
        return redirect("financeiro:lancamento", pk=pk)
    if l.estornos.exists() or l.medicao_id or l.conta_origem_id:
        messages.error(request, "Não dá para excluir: tem estorno ligado, ou nasceu de uma medição ou de uma conta.")
        return redirect("financeiro:lancamento", pk=pk)
    l.delete()
    messages.success(request, "Lançamento excluído. Fica registrado na auditoria.")
    return redirect("financeiro:lancamentos")


@so_financeiro
@com_obra
def exportar_lancamentos(request, obra):
    """Os lançamentos filtrados em .xlsx, para o contador ou para conferência."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    qs = consultas.lancamentos_filtrados(obra, request.GET)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    ws.append(["Data", "Mês", "Categoria", "Tipo", "Prestador / fornecedor", "Descrição", "NF", "Qtd", "Un",
               "Forma", "Conta", "Responsável", "Valor", "Comprovante"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0071E3")
    for l in qs:
        ws.append([l.data, l.competencia.strftime("%m/%Y"), l.categoria.nome,
                   l.subcategoria.nome if l.subcategoria_id else l.get_tipo_mao_de_obra_display(),
                   l.favorecido, l.descricao, l.numero_nf, float(l.quantidade) if l.quantidade else None, l.unidade,
                   l.get_forma_pagamento_display(), str(l.conta_bancaria or ""), str(l.responsavel or ""),
                   float(l.valor), "sim" if l.comprovante else "não"])
    for linha in ws.iter_rows(min_row=2):
        linha[0].number_format = "DD/MM/YYYY"
        linha[12].number_format = '"R$" #,##0.00'
    for col, largura in zip("ABCDEFGHIJKLMN", (11, 8, 20, 18, 28, 40, 10, 8, 6, 14, 10, 14, 14, 12)):
        ws.column_dimensions[col].width = largura
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="lancamentos-{date.today():%Y-%m-%d}.xlsx"'
    return resp


@so_financeiro
@com_obra
def fechamento(request, obra, ano=None, mes=None):
    """O bloco do mês da planilha (MÊS 06 JUNHO...), montado pelo sistema."""
    hoje = date.today()
    try:
        competencia = date(ano, mes, 1) if ano else hoje.replace(day=1)
    except ValueError:
        raise Http404
    blocos, total = consultas.fechamento_do_mes(obra, competencia)
    anterior = (competencia - timedelta(days=1)).replace(day=1)
    return render(request, "financeiro/fechamento.html", _ctx(
        request, obra, blocos=blocos, total=total, competencia=competencia,
        titulo_mes=consultas.nome_mes(competencia), anterior=anterior, seguinte=consultas.proximo_mes(competencia),
        hoje=hoje,
    ))


@so_financeiro
@com_obra
def pendencias(request, obra):
    return render(request, "financeiro/pendencias.html", _ctx(request, obra, itens=consultas.pendencias(obra)))


# ---------------------------------------------------------------------------
# mão de obra
# ---------------------------------------------------------------------------

@so_financeiro
@com_obra
def mao_de_obra(request, obra):
    lista = list(Prestador.objects.filter(lancamentos__obra=obra, lancamentos__categoria__codigo=Categoria.MAO_DE_OBRA)
                 .annotate(total=Coalesce(Sum("lancamentos__valor"), ZERO), qtd=Count("lancamentos"),
                           diarias=Coalesce(Sum("lancamentos__diarias"), ZERO), ultimo=Max("lancamentos__data"))
                 .order_by("-total"))
    for p in lista:
        p.vales = regras.saldo_de_vales(p, obra)
    hoje = date.today()
    inicio_semana = hoje - timedelta(days=hoje.weekday())
    semana = obra.lancamentos.filter(categoria__codigo=Categoria.MAO_DE_OBRA, data__gte=inicio_semana)
    return render(request, "financeiro/mao_de_obra.html", _ctx(
        request, obra, lista=lista, total=sum((p.total for p in lista), ZERO),
        sem_pagamento=Prestador.objects.filter(ativo=True).exclude(pk__in=[p.pk for p in lista]),
        semana_total=semana.aggregate(t=Coalesce(Sum("valor"), ZERO))["t"],
        semana_diarias=semana.aggregate(t=Coalesce(Sum("diarias"), ZERO))["t"],
        semana_pessoas=semana.values("prestador").distinct().count(),
        vales_total=sum((p.vales for p in lista), ZERO), inicio_semana=inicio_semana,
    ))


@exige_papel(Perfil.JURIDICO)
def prestador_detalhe(request, pk):
    """O 'acerto de mão de obra' (Planilha6): pagamentos mês a mês, vales em aberto."""
    p = get_object_or_404(Prestador, pk=pk)
    obra = obra_atual(request)
    ve_valores = papel_de(request.user) == Perfil.SOCIO
    meses, total, vales = [], ZERO, ZERO
    if ve_valores and obra:
        lancs = list(p.lancamentos.filter(obra=obra).select_related("categoria").order_by("data", "id"))
        por_mes = {}
        for l in lancs:
            por_mes.setdefault(l.competencia, []).append(l)
        meses = [{"mes": consultas.nome_mes(m), "itens": itens, "total": sum((i.valor for i in itens), ZERO),
                  "diarias": sum((i.diarias for i in itens), ZERO)} for m, itens in sorted(por_mes.items(), reverse=True)]
        total = sum((l.valor for l in lancs), ZERO)
        vales = regras.saldo_de_vales(p, obra)
    return render(request, "financeiro/prestador_detalhe.html", _ctx(
        request, obra, p=p, meses=meses, total=total, vales=vales, ve_valores=ve_valores,
        empreitas=p.empreitas.all()))


# ---------------------------------------------------------------------------
# empreitas
# ---------------------------------------------------------------------------

@so_financeiro
@com_obra
def empreitas(request, obra):
    lista = list(obra.empreitas.select_related("fornecedor", "prestador"))
    for e in lista:
        e.r = regras.resumo_empreita(e)
    return render(request, "financeiro/empreitas.html", _ctx(request, obra, lista=lista))


@so_financeiro
@com_obra
def empreita_detalhe(request, obra, pk):
    e = get_object_or_404(Empreita, pk=pk, obra=obra)
    return render(request, "financeiro/empreita_detalhe.html", _ctx(
        request, obra, e=e, r=regras.resumo_empreita(e),
        pagamentos=e.pagamentos.select_related("categoria").order_by("data"),
        categoria_pagamento=Categoria.MAO_DE_OBRA if e.prestador_id and not e.fornecedor_id else Categoria.SERVICOS,
    ))


@so_financeiro
@com_obra
def empreita_form(request, obra, pk=None):
    e = get_object_or_404(Empreita, pk=pk, obra=obra) if pk else None
    form = EmpreitaForm(request.POST or None, request.FILES or None, instance=e, obra=obra)
    if request.method == "POST" and form.is_valid():
        obj = form.save()
        messages.success(request, "Empreita salva.")
        return redirect("financeiro:empreita", pk=obj.pk)
    return render(request, "financeiro/form_simples.html", _ctx(
        request, obra, form=form, titulo="Editar empreita" if e else "Nova empreita",
        subtitulo="Serviço contratado por preço global ou por medida (m², metro...).",
        voltar=reverse("financeiro:empreitas"), fornecedores=Fornecedor.objects.filter(ativo=True)))


# ---------------------------------------------------------------------------
# contas a pagar
# ---------------------------------------------------------------------------

@so_financeiro
def contas(request):
    regras.gerar_todas_as_locacoes()
    obra = obra_atual(request)
    hoje = date.today()
    ver = request.GET.get("ver", "abertas")
    todas = request.GET.get("escopo") == "todas"
    escopo = None if todas else obra
    qs = ContaAPagar.objects.select_related("fornecedor", "obra", "lancamento", "locacao")
    if escopo:
        qs = qs.filter(obra=escopo)
    if ver == "pagas":
        qs = qs.filter(pago_em__isnull=False).order_by("-pago_em")
    elif ver == "vencidas":
        qs = qs.filter(pago_em__isnull=True, vencimento__lt=hoje)
    elif ver == "semana":
        qs = qs.filter(pago_em__isnull=True, vencimento__gte=hoje, vencimento__lte=hoje + timedelta(days=7))
    else:
        qs = qs.filter(pago_em__isnull=True)
    itens = list(qs[:300])
    for c in itens:
        c.sit = c.situacao(hoje)
    return render(request, "financeiro/contas.html", {
        "obra": obra, "obras": Obra.objects.all(), "itens": itens, "ver": ver, "todas": todas, "hoje": hoje,
        "resumo": consultas.contas_resumo(escopo, hoje), "form_pagar": PagarContaForm(initial={"pago_em": hoje}),
        "locacoes": LocacaoRecorrente.objects.filter(obra=obra).select_related("fornecedor") if obra else [],
        "total_lista": sum((c.valor for c in itens), ZERO),
    })


@so_financeiro
@require_POST
def conta_pagar(request, pk):
    conta = get_object_or_404(ContaAPagar, pk=pk)
    form = PagarContaForm(request.POST, request.FILES)
    if form.is_valid():
        regras.pagar_conta(conta, form.cleaned_data["pago_em"], form.cleaned_data.get("conta_bancaria"),
                           form.cleaned_data.get("comprovante"), usuario=request.user)
        messages.success(request, f"{conta.fornecedor} — R$ {conta.valor} paga.")
    else:
        messages.error(request, "Informe a data do pagamento.")
    return redirect(_voltar(request, "financeiro:contas"))


@so_financeiro
@require_POST
def conta_desfazer(request, pk):
    regras.desfazer_pagamento(get_object_or_404(ContaAPagar, pk=pk))
    messages.success(request, "Pagamento desfeito.")
    return redirect(_voltar(request, "financeiro:contas"))


def _voltar(request, padrao):
    destino = request.POST.get("voltar") or ""
    return destino if destino.startswith("/") and not destino.startswith("//") else reverse(padrao)


@so_financeiro
def conta_form(request, pk=None):
    obra = obra_atual(request)
    conta = get_object_or_404(ContaAPagar, pk=pk) if pk else None
    if conta and conta.lancamento_id:
        messages.info(request, "Esta conta é parcela de uma compra: edite a compra.")
        return redirect("financeiro:lancamento", pk=conta.lancamento_id)
    if conta and conta.locacao_id:
        messages.info(request, "Esta conta é de uma locação mensal: edite a locação.")
        return redirect("financeiro:locacao_editar", pk=conta.locacao_id)
    form = ContaForm(request.POST or None, request.FILES or None, instance=conta,
                     initial=None if conta else {"obra": obra})
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Conta salva.")
        return redirect("financeiro:contas")
    return render(request, "financeiro/form_simples.html", _ctx(
        request, obra, form=form, titulo="Editar conta" if conta else "Nova conta avulsa",
        subtitulo="Para compra a prazo, lance a compra: as parcelas viram contas sozinhas.",
        voltar=reverse("financeiro:contas"), fornecedores=Fornecedor.objects.filter(ativo=True)))


@so_financeiro
@com_obra
def locacao_form(request, obra, pk=None):
    loc = get_object_or_404(LocacaoRecorrente, pk=pk) if pk else None
    form = LocacaoForm(request.POST or None, instance=loc)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        if not loc:
            obj.obra = obra
        obj.save()
        if loc:  # mudou valor ou fim: corrige as contas ainda não pagas
            obj.contas.filter(pago_em__isnull=True).update(valor=obj.valor_mensal)
            if obj.fim:
                obj.contas.filter(pago_em__isnull=True, vencimento__gt=obj.fim).delete()
            if not obj.ativa:
                obj.contas.filter(pago_em__isnull=True).delete()
        n = regras.gerar_contas_de_locacao(obj)
        messages.success(request, "Locação salva." + (f" {n} conta(s) gerada(s)." if n else ""))
        return redirect("financeiro:contas")
    return render(request, "financeiro/form_simples.html", _ctx(
        request, obra, form=form, titulo="Editar locação mensal" if loc else "Nova locação mensal",
        subtitulo="Contêiner, kit, gerador: a conta de cada mês aparece sozinha. Paga, vira gasto da obra.",
        voltar=reverse("financeiro:contas"), fornecedores=Fornecedor.objects.filter(ativo=True)))


# ---------------------------------------------------------------------------
# contrato e medições
# ---------------------------------------------------------------------------

@so_financeiro
@com_obra
def contrato(request, obra):
    hoje = date.today()
    medicoes = list(obra.medicoes.all())
    for m in medicoes:
        m.sit = m.situacao(hoje)
        m.imposto = m.impostos.first()
    r = consultas.resumo_obra(obra, hoje)
    return render(request, "financeiro/contrato.html", _ctx(
        request, obra, r=r, medicoes=medicoes, aditivos=obra.aditivos.all(),
        soma_pct=sum((m.percentual or ZERO for m in medicoes), ZERO),
        faltando_pct=consultas.pct(r["a_faturar"], r["contrato"])))


@so_financeiro
@com_obra
def medicao_form(request, obra, pk=None):
    med = get_object_or_404(Medicao, pk=pk, obra=obra) if pk else None
    inicial = None
    if not med:
        n = obra.medicoes.count()
        inicial = {"descricao": "Entrada" if n == 0 else f"Medição {n:02d}"}
    form = MedicaoForm(request.POST or None, request.FILES or None, instance=med, obra=obra, initial=inicial)
    if request.method == "POST" and form.is_valid():
        with transaction.atomic():
            obj = form.save()
            imposto = regras.sincronizar_imposto(obj)
        msg = "Medição salva."
        if imposto:
            msg += f" Imposto de R$ {imposto.valor} lançado automaticamente."
        messages.success(request, msg)
        return redirect("financeiro:contrato")
    return render(request, "financeiro/form_simples.html", _ctx(
        request, obra, form=form, titulo="Editar medição" if med else "Nova medição / NF de entrada",
        subtitulo=f"Contrato com aditivos: R$ {obra.valor_total:,.2f}. Imposto da obra: {obra.aliquota_imposto}%."
                  .replace(",", "X").replace(".", ",").replace("X", "."),
        voltar=reverse("financeiro:contrato"),
        excluir=reverse("financeiro:medicao_excluir", args=[med.pk]) if med else None))


@so_financeiro
@require_POST
def medicao_excluir(request, pk):
    get_object_or_404(Medicao, pk=pk).delete()
    messages.success(request, "Medição excluída, junto com o imposto que ela gerou.")
    return redirect("financeiro:contrato")


@so_financeiro
@com_obra
def aditivo_form(request, obra, pk=None):
    adt = get_object_or_404(Aditivo, pk=pk, obra=obra) if pk else None
    form = AditivoForm(request.POST or None, instance=adt)
    if request.method == "POST" and form.is_valid():
        obj = form.save(commit=False)
        obj.obra = obra
        obj.save()
        messages.success(request, "Aditivo salvo. O valor do contrato foi atualizado.")
        return redirect("financeiro:contrato")
    return render(request, "financeiro/form_simples.html", _ctx(
        request, obra, form=form, titulo="Editar aditivo" if adt else "Novo aditivo",
        voltar=reverse("financeiro:contrato")))


# ---------------------------------------------------------------------------
# obras
# ---------------------------------------------------------------------------

@so_financeiro
def obras(request):
    return render(request, "financeiro/obras.html", {
        "v": consultas.visao_geral(), "obra": obra_atual(request), "obras": Obra.objects.all()})


@so_financeiro
def obra_form(request, pk=None):
    obra = get_object_or_404(Obra, pk=pk) if pk else None
    form = ObraForm(request.POST or None, instance=obra)
    if request.method == "POST" and form.is_valid():
        obj = form.save()
        n = regras.criar_checklist(obj)
        request.session["obra_id"] = obj.pk
        messages.success(request, "Obra salva." + (f" Checklist com {n} itens criado." if n else ""))
        return redirect("financeiro:painel")
    return render(request, "financeiro/form_simples.html", {
        "form": form, "obra": obra or obra_atual(request), "obras": Obra.objects.all(),
        "titulo": f"Editar {obra.nome}" if obra else "Nova obra",
        "subtitulo": "Valor, prazo, imposto e retenção valem para todos os cálculos da obra.",
        "voltar": reverse("financeiro:obras"),
    })


# ---------------------------------------------------------------------------
# relatórios
# ---------------------------------------------------------------------------

@so_financeiro
@com_obra
def relatorios(request, obra):
    return render(request, "financeiro/relatorios.html", _ctx(
        request, obra,
        geral=consultas.matriz_mes_categoria(obra),
        materiais=consultas.matriz_materiais(obra),
        concreto=consultas.controle_por_quantidade(obra, "Concreto"),
        diesel=consultas.controle_por_quantidade(obra, "Diesel", Categoria.COMBUSTIVEL),
        gasolina=consultas.controle_por_quantidade(obra, "Gasolina", Categoria.COMBUSTIVEL),
        aba=request.GET.get("aba", "geral"),
    ))


# ---------------------------------------------------------------------------
# checklist
# ---------------------------------------------------------------------------

@exige_papel(Perfil.OPERACAO, Perfil.JURIDICO, Perfil.COMPRAS)
@com_obra
def checklist(request, obra):
    regras.criar_checklist(obra)
    grupos = {}
    for item in obra.checklist.select_related("atualizado_por"):
        grupos.setdefault(item.grupo, []).append(item)
    feitos = obra.checklist.filter(situacao__in=[ItemChecklist.FEITO, ItemChecklist.NAO_APLICA]).count()
    total = obra.checklist.count()
    return render(request, "financeiro/checklist.html", _ctx(
        request, obra, grupos=grupos, feitos=feitos, total=total, pct=consultas.pct(feitos, total),
        situacoes=ItemChecklist.SITUACOES, hoje=date.today()))


@exige_papel(Perfil.OPERACAO, Perfil.JURIDICO, Perfil.COMPRAS)
@require_POST
def checklist_item(request, pk):
    item = get_object_or_404(ItemChecklist, pk=pk)
    dados = request.POST.copy()
    for campo in ("situacao", "prazo", "link"):  # o HTMX manda só o campo que mudou
        if campo not in dados:
            valor = getattr(item, campo)
            dados[campo] = valor.isoformat() if hasattr(valor, "isoformat") else (valor or "")
    form = ItemChecklistForm(dados, request.FILES, instance=item)
    if form.is_valid():
        item = form.save(commit=False)
        item.atualizado_por = request.user
        item.save()
    if request.headers.get("HX-Request"):
        return render(request, "financeiro/_checklist_item.html",
                      {"item": item, "situacoes": ItemChecklist.SITUACOES, "hoje": date.today()})
    return redirect("financeiro:checklist")


# ---------------------------------------------------------------------------
# arquivos (comprovante, NF, boleto): só com login e permissão
# ---------------------------------------------------------------------------

MODELOS_ARQUIVO = {
    "lancamento": (Lancamento, ("comprovante", "arquivo_nf")),
    "medicao": (Medicao, ("arquivo_nf",)),
    "conta": (ContaAPagar, ("boleto", "comprovante")),
    "empreita": (Empreita, ("contrato",)),
    "checklist": (ItemChecklist, ("arquivo",)),
}


@exige_papel(*PODE_LANCAR, Perfil.OPERACAO, Perfil.JURIDICO)
def arquivo(request, modelo, pk, campo):
    if modelo not in MODELOS_ARQUIVO or campo not in MODELOS_ARQUIVO[modelo][1]:
        raise Http404
    obj = get_object_or_404(MODELOS_ARQUIVO[modelo][0], pk=pk)
    papel = papel_de(request.user)
    if modelo == "lancamento" and not _pode_categoria(request, obj.categoria.codigo):
        raise PermissionDenied
    if modelo in ("medicao", "conta", "empreita") and papel != Perfil.SOCIO:
        raise PermissionDenied
    f = getattr(obj, campo)
    if not f:
        raise Http404
    return FileResponse(f.open("rb"), as_attachment=False, filename=f.name.rsplit("/", 1)[-1])


# ---------------------------------------------------------------------------
# responsáveis (quem compra / paga)
# ---------------------------------------------------------------------------

@so_financeiro
def responsaveis(request):
    if request.method == "POST":
        nome = (request.POST.get("nome") or "").strip()
        if nome:
            Responsavel.objects.get_or_create(nome=nome.title())
            messages.success(request, f"{nome.title()} cadastrado.")
        return redirect("financeiro:responsaveis")
    return render(request, "financeiro/responsaveis.html", {
        "obra": obra_atual(request), "obras": Obra.objects.all(),
        "lista": Responsavel.objects.annotate(qtd=Count("lancamento"), total=Coalesce(Sum("lancamento__valor"), ZERO)),
    })
